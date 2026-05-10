from __future__ import annotations

import io
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from ..schemas.docx import BulletLine

# ── Renk paleti ───────────────────────────────────────────────────────────────
_COLOR_TITLE    = RGBColor(0x1F, 0x39, 0x64)
_COLOR_HEADER   = RGBColor(0x1F, 0x39, 0x64)   # ❖ başlık
_COLOR_PROJECT  = RGBColor(0x2E, 0x75, 0xB6)   # • proje adı (grup altında)
_COLOR_BULLET   = RGBColor(0x00, 0x00, 0x00)   # ▪ yorum
_COLOR_FOLLOWUP = RGBColor(0xC0, 0x6B, 0x00)   # ▪ takip


def _tight(p) -> None:
    pPr = p._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "40")
    spacing.set(qn("w:after"), "40")
    pPr.append(spacing)


def _clean(label: str) -> str:
    """Köşeli parantezleri koru — [Proje Adı] olarak bırak."""
    return label.strip()


def _add_header(doc: Document, text: str) -> None:
    """❖ Grup veya bağımsız proje başlığı."""
    p = doc.add_paragraph()
    _tight(p)
    p.paragraph_format.left_indent  = Inches(0)
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run(f"\u2756 {text}")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = _COLOR_HEADER


def _add_project(doc: Document, text: str, indent: float) -> None:
    """• Proje adı — grup altında."""
    p = doc.add_paragraph()
    _tight(p)
    p.paragraph_format.left_indent = Inches(indent)
    run = p.add_run(f"\u2022 {text}")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = _COLOR_PROJECT


def _add_bullet(doc: Document, text: str, indent: float, followup: bool = False) -> None:
    """▪ Yorum / takip satırı."""
    p = doc.add_paragraph()
    _tight(p)
    p.paragraph_format.left_indent = Inches(indent)
    label = f"[Takip] {text}" if followup else text
    run = p.add_run(f"\u25aa {label}")
    run.font.size = Pt(10)
    run.font.color.rgb = _COLOR_FOLLOWUP if followup else _COLOR_BULLET


def build_docx(lines: list[BulletLine], title: str) -> bytes:
    doc = Document()

    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.25)
        section.right_margin  = Inches(1.25)

    # Belge başlığı
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = _COLOR_TITLE
    doc.add_paragraph()

    inside_group = False

    for line in lines:
        b0    = line.bullet0
        label = b0.strip() if b0 else ""

        # DURUM 1: bullet0 var, bullet1+bullet2 yok → saf grup başlığı
        # { bullet0: "[Grup]", bullet1: null, bullet2: null }
        if label and not line.bullet1 and not line.bullet2:
            inside_group = True
            _add_header(doc, _clean(label))
            continue

        # DURUM 2: bullet0 var, bullet1 var, bullet2 yok
        # → bağımsız proje: ❖[Proje] + ▪ bullet1 yorumları
        if label and line.bullet1 and not line.bullet2:
            inside_group = False
            _add_header(doc, _clean(label))
            for item in line.bullet1:
                _add_bullet(doc, item, indent=0.35)
            for item in (line.bullet3 or []):
                _add_bullet(doc, item, indent=0.35, followup=True)
            continue

        # DURUM 3: bullet0 var, bullet2 de var
        # → bağımsız proje: ❖[Proje] + ▪ bullet2 yorumları
        if label and line.bullet2:
            inside_group = False
            _add_header(doc, _clean(label))
            for item in line.bullet2:
                _add_bullet(doc, item, indent=0.35)
            for item in (line.bullet3 or []):
                _add_bullet(doc, item, indent=0.35, followup=True)
            continue

        # DURUM 4: bullet0 = None → grup içindeki proje satırı
        # { bullet0: null, bullet1: ["Proje Adı"], bullet2: ["yorum1", ...] }
        if not label:
            proj_indent   = 0.30
            bullet_indent = 0.65

            if line.bullet1:
                proj_name = line.bullet1[0] if isinstance(line.bullet1, list) else str(line.bullet1)
                _add_project(doc, proj_name, indent=proj_indent)
                for item in (line.bullet1[1:] if isinstance(line.bullet1, list) else []):
                    _add_bullet(doc, item, indent=bullet_indent)

            for item in (line.bullet2 or []):
                _add_bullet(doc, item, indent=bullet_indent)

            for item in (line.bullet3 or []):
                _add_bullet(doc, item, indent=bullet_indent, followup=True)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()